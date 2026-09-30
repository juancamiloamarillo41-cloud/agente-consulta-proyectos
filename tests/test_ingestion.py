import docx
import pytest

from project_agent.ingestion import load_document, split_into_chunks

EXPECTED_CODES = {"PC-2025-014", "PC-2025-027", "PC-2025-033", "PC-2026-006"}


def _section(chunks, code, prefix):
    return next(c for c in chunks if c.project_code == code and c.section.startswith(prefix))


def test_loads_all_four_reports_with_codes_and_clients(documents):
    assert set(documents) == EXPECTED_CODES
    assert documents["PC-2025-027"].client == "Plásticos del Pacífico S.A."
    assert documents["PC-2025-033"].title == "Reducción de tiempos de admisión y espera en consulta externa"


def test_page_footers_are_removed(documents):
    for doc in documents.values():
        assert "Documento de uso interno" not in doc.full_text


@pytest.mark.parametrize(
    "code, expected_sections",
    [
        ("PC-2025-014", ["Datos generales del proyecto", "1. Resumen ejecutivo", "5. Resultados", "7. Recomendaciones y próximos pasos"]),
        ("PC-2025-027", ["4. Alcance", "7. Resultados", "Anexo A. Principales causas de parada en la línea base (Línea 1)"]),
        ("PC-2026-006", ["3. Objetivos", "6. Resultados", "7. Lecciones aprendidas"]),
    ],
)
def test_splits_by_top_level_sections(chunks, code, expected_sections):
    sections = [c.section for c in chunks if c.project_code == code]
    for expected in expected_sections:
        assert expected in sections


def test_subsections_stay_with_their_section_context(chunks):
    # La advertencia "exclusivamente Línea 1" debe viajar junto a la tabla de OEE.
    results = _section(chunks, "PC-2025-027", "7. Resultados")
    assert "exclusivamente a la Línea 1" in results.text
    assert "| OEE | 58% | 71% | +13 pp |" in results.text


def test_tables_keep_row_column_relationship(chunks):
    coop = _section(chunks, "PC-2025-014", "5. Resultados")
    assert "| Tasa de abandono de solicitudes | % del total | 18% | ≤ 10% | 11% | -7 pp | No |" in coop.text

    canasta = _section(chunks, "PC-2026-006", "6. Resultados")
    assert "| Días de inventario en tienda | 38 días | ≤ 30 días | 31 días | Parcialmente cumplido |" in canasta.text


def test_word_exported_pdf_tables_are_cleaned(chunks):
    clinic = _section(chunks, "PC-2025-033", "6. Resultados")
    assert "| Tiempo total de espera del paciente | 52 min | reducción ≥ 20% | 39,5 min | -24% |" in clinic.text
    assert "| Satisfacción del paciente (NPS) | 18 | sin meta | 37 | +19 puntos |" in clinic.text
    # La nota que invalida el 30% preliminar queda en el mismo fragmento.
    assert "constituye el dato oficial de cierre" in clinic.text

    front = _section(chunks, "PC-2025-033", "Datos generales")
    assert "| Cliente | Clínica Santa Lucía del Valle, clínica privada de especialidades con 38 consultorios y 60 camas |" in front.text


def test_chunks_carry_citation_metadata(chunks):
    chunk = _section(chunks, "PC-2025-033", "6. Resultados")
    assert chunk.pages == (2, 3)
    assert chunk.citation == "Informe_Cierre_PC-2025-033_Clinica_Santa_Lucia.pdf — 6. Resultados, págs. 2-3"
    assert chunk.text.startswith("[PC-2025-033 · Clínica Santa Lucía del Valle · ")


def test_docx_loader(tmp_path):
    # El enunciado indica que uno de los informes llega en Word: se valida el lector con un .docx sintético.
    document = docx.Document()
    document.add_paragraph("Proyecto de prueba", style="Title")
    document.add_paragraph("Empresa Demo S.A.", style="Subtitle")
    document.add_paragraph("Código de proyecto PC-2099-001")
    document.add_heading("1. Resumen ejecutivo", level=1)
    document.add_paragraph("Se redujo el tiempo de ciclo.")
    document.add_heading("2. Resultados", level=1)
    table = document.add_table(rows=3, cols=3)
    for i, row in enumerate([["Indicador", "Línea base", "Resultado"], ["Tiempo de ciclo", "10 días", "6 días"]]):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    merged = table.cell(2, 0).merge(table.cell(2, 2))
    merged.text = "Nota: fila combinada"
    document.add_heading("2.1 Detalle", level=2)
    document.add_paragraph("Texto posterior a la tabla.")
    path = tmp_path / "Informe_Demo.docx"
    document.save(path)

    doc = load_document(path)
    assert doc.project_code == "PC-2099-001"
    assert doc.title == "Proyecto de prueba"
    assert doc.client == "Empresa Demo S.A."

    chunks = split_into_chunks(doc)
    assert [c.section for c in chunks] == ["Datos generales del proyecto", "1. Resumen ejecutivo", "2. Resultados"]
    results = chunks[2]
    assert results.pages == ()
    assert "| Tiempo de ciclo | 10 días | 6 días |" in results.text
    assert "| Nota: fila combinada |" in results.text
    assert results.text.index("| Tiempo de ciclo") < results.text.index("### 2.1 Detalle")


def test_unsupported_format_is_rejected(tmp_path):
    path = tmp_path / "informe.txt"
    path.write_text("hola", encoding="utf-8")
    with pytest.raises(ValueError, match="Formato no soportado"):
        load_document(path)
